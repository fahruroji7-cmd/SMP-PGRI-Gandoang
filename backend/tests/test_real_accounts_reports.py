import io
import os

import openpyxl
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")


def login(email, password):
    session = requests.Session()
    response = session.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password}, timeout=20)
    return session, response


@pytest.mark.parametrize("email,password,role", [
    ("admin@absenspg.local", "Admin123!", "Admin"),
    ("guru@absenspg.local", "Guru123!", "Guru"),
])
def test_login_me_permissions_and_cookie(email, password, role):
    session, response = login(email, password)
    assert response.status_code == 200
    assert response.json()["role"] == role
    cookie = response.cookies.get("access_token")
    assert cookie and response.cookies.get_dict()
    assert any(c.name == "access_token" and c.has_nonstandard_attr("HttpOnly") for c in response.cookies)
    me = session.get(f"{BASE_URL}/api/auth/me", timeout=20)
    assert me.status_code == 200 and me.json()["email"] == email
    permissions = session.get(f"{BASE_URL}/api/auth/permissions", timeout=20)
    assert permissions.status_code == 200
    assert permissions.json()["can_manage_master_data"] is (role == "Admin")


@pytest.mark.parametrize("kind", ["attendance", "grades", "journals"])
def test_xlsx_export_is_valid(kind):
    session, response = login("admin@absenspg.local", "Admin123!")
    assert response.status_code == 200
    exported = session.get(f"{BASE_URL}/api/reports/export", params={"kind": kind}, timeout=20)
    assert exported.status_code == 200
    assert "spreadsheetml" in exported.headers.get("content-type", "")
    workbook = openpyxl.load_workbook(io.BytesIO(exported.content), read_only=True)
    assert workbook.active.max_row >= 1
    assert workbook.active.cell(1, 1).value == "No"


@pytest.mark.parametrize("kind", ["attendance", "grades", "journals"])
def test_print_report_is_standalone_html(kind):
    session, response = login("admin@absenspg.local", "Admin123!")
    assert response.status_code == 200
    printed = session.get(f"{BASE_URL}/api/reports/print", params={"kind": kind}, timeout=20)
    assert printed.status_code == 200
    assert "<table>" in printed.text and "window.print()" in printed.text


def test_logout_invalidates_cookie_client():
    session, response = login("guru@absenspg.local", "Guru123!")
    assert response.status_code == 200
    assert session.get(f"{BASE_URL}/api/auth/me", timeout=20).status_code == 200
    logged_out = session.post(f"{BASE_URL}/api/auth/logout", timeout=20)
    assert logged_out.status_code == 200
    assert session.get(f"{BASE_URL}/api/auth/me", timeout=20).status_code == 401