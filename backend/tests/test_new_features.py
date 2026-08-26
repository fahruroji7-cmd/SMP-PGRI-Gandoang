"""Backend tests for iteration 8 features: schedules, master edit/delete, dashboard stats, report filters."""
import io
import os
import re
from pathlib import Path

import pytest
import requests
from dotenv import dotenv_values
from openpyxl import load_workbook

frontend_env = dotenv_values("/app/frontend/.env")
base_url = os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")
if not base_url:
    raise RuntimeError("REACT_APP_BACKEND_URL missing")
BASE_URL = base_url.rstrip("/")


def creds(role):
    content = Path("/app/memory/test_credentials.md").read_text(encoding="utf-8")
    section = content.split(f"## {role}")[1]
    email = re.search(r"Email:\s*`([^`]+)`", section).group(1)
    password = re.search(r"Password:\s*`([^`]+)`", section).group(1)
    return email, password


def login(role):
    s = requests.Session()
    email, password = creds(role)
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login {role} failed: {r.status_code} {r.text[:300]}"
    assert r.json()["role"] == role
    assert "access_token" in s.cookies, "httpOnly access_token cookie not set"
    return s


@pytest.fixture(scope="module")
def admin():
    return login("Admin")


@pytest.fixture(scope="module")
def guru():
    return login("Guru")


# ---------- Auth regression ----------
class TestAuth:
    def test_admin_and_guru_login_and_me(self, admin, guru):
        for s, role in ((admin, "Admin"), (guru, "Guru")):
            r = s.get(f"{BASE_URL}/api/auth/me", timeout=30)
            assert r.status_code == 200
            assert r.json()["role"] == role

    def test_bad_password_401(self):
        r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": creds("Admin")[0], "password": "wrong-pass"}, timeout=30)
        assert r.status_code == 401

    def test_bcrypt_hash_format(self):
        import asyncio

        from motor.motor_asyncio import AsyncIOMotorClient
        env = dotenv_values("/app/backend/.env")

        async def check():
            cl = AsyncIOMotorClient(env["MONGO_URL"])
            u = await cl[env["DB_NAME"]].users.find_one({"email": creds("Admin")[0]})
            cl.close()
            return u
        u = asyncio.get_event_loop().run_until_complete(check())
        assert u and u["password_hash"].startswith("$2b$"), u and u["password_hash"][:6]

    def test_unauthenticated_blocked(self):
        for path in ["/api/dashboard/stats", "/api/schedules", "/api/masters", "/api/reports/rows"]:
            r = requests.get(f"{BASE_URL}{path}", timeout=30)
            assert r.status_code == 401, f"{path} -> {r.status_code}"


# ---------- Schedules ----------
class TestSchedules:
    created = []

    def test_create_list_persist_delete(self, guru):
        payload = {"day": "Rabu", "start_time": "07:30", "end_time": "08:50", "class_name": "7A", "subject": "Matematika", "room": "TEST_R1"}
        r = guru.post(f"{BASE_URL}/api/schedules", json=payload, timeout=30)
        assert r.status_code == 200, r.text[:300]
        doc = r.json()
        assert doc["id"] and doc["day"] == "Rabu" and doc["room"] == "TEST_R1"
        assert "_id" not in doc

        lst = guru.get(f"{BASE_URL}/api/schedules", timeout=30)
        assert lst.status_code == 200
        rows = lst.json()
        assert any(x["id"] == doc["id"] for x in rows), "created schedule not persisted"
        assert all("_id" not in x for x in rows)

        d = guru.delete(f"{BASE_URL}/api/schedules/{doc['id']}", timeout=30)
        assert d.status_code == 200
        rows2 = guru.get(f"{BASE_URL}/api/schedules", timeout=30).json()
        assert not any(x["id"] == doc["id"] for x in rows2), "schedule not deleted"

    def test_delete_unknown_404(self, guru):
        assert guru.delete(f"{BASE_URL}/api/schedules/does-not-exist", timeout=30).status_code == 404

    def test_invalid_day_422(self, guru):
        r = guru.post(f"{BASE_URL}/api/schedules", json={"day": "Minggu", "start_time": "07:00", "end_time": "08:00", "class_name": "7A", "subject": "IPA"}, timeout=30)
        assert r.status_code == 422, r.status_code

    def test_sorted_by_day_then_time(self, guru):
        made = []
        for day, st in [("Selasa", "09:00"), ("Senin", "10:00"), ("Senin", "06:30")]:
            r = guru.post(f"{BASE_URL}/api/schedules", json={"day": day, "start_time": st, "end_time": "23:00", "class_name": "7A", "subject": "IPA", "room": "TEST_SORT"}, timeout=30)
            assert r.status_code == 200
            made.append(r.json()["id"])
        rows = [x for x in guru.get(f"{BASE_URL}/api/schedules", timeout=30).json() if x.get("room") == "TEST_SORT"]
        keys = [(x["day"], x["start_time"]) for x in rows]
        order = {"Senin": 0, "Selasa": 1}
        assert keys == sorted(keys, key=lambda k: (order[k[0]], k[1])), keys
        for i in made:
            guru.delete(f"{BASE_URL}/api/schedules/{i}", timeout=30)


# ---------- Master edit/delete + RBAC ----------
class TestMasterCrud:
    @pytest.mark.parametrize("resource,payload,updated", [
        ("teachers", {"name": "TEST_Guru A"}, {"name": "TEST_Guru A Edited"}),
        ("classes", {"name": "TEST_9Z"}, {"name": "TEST_9Y"}),
        ("subjects", {"name": "TEST_Seni"}, {"name": "TEST_Seni Musik"}),
    ])
    def test_admin_crud_cycle(self, admin, resource, payload, updated):
        c = admin.post(f"{BASE_URL}/api/{resource}", json=payload, timeout=30)
        assert c.status_code == 200, c.text[:300]
        item_id = c.json()["id"]
        u = admin.put(f"{BASE_URL}/api/{resource}/{item_id}", json=updated, timeout=30)
        assert u.status_code == 200, u.text[:300]
        masters = admin.get(f"{BASE_URL}/api/masters", timeout=30).json()
        rows = masters[resource]
        row = next((x for x in rows if x["id"] == item_id), None)
        assert row and row["name"] == updated["name"], "update not persisted"
        d = admin.delete(f"{BASE_URL}/api/{resource}/{item_id}", timeout=30)
        assert d.status_code == 200
        rows2 = admin.get(f"{BASE_URL}/api/masters", timeout=30).json()[resource]
        assert not any(x["id"] == item_id for x in rows2), "delete not persisted"

    def test_admin_student_crud_cycle(self, admin):
        c = admin.post(f"{BASE_URL}/api/students", json={"name": "TEST_Siswa Satu", "class_name": "7A"}, timeout=30)
        assert c.status_code == 200
        sid = c.json()["id"]
        u = admin.put(f"{BASE_URL}/api/students/{sid}", json={"name": "TEST_Siswa Edited", "class_name": "7B"}, timeout=30)
        assert u.status_code == 200
        got = [x for x in admin.get(f"{BASE_URL}/api/students", params={"class_name": "7B"}, timeout=30).json() if x["id"] == sid]
        assert got and got[0]["name"] == "TEST_Siswa Edited" and got[0]["class_name"] == "7B"
        assert admin.delete(f"{BASE_URL}/api/students/{sid}", timeout=30).status_code == 200
        assert not any(x["id"] == sid for x in admin.get(f"{BASE_URL}/api/students", timeout=30).json())

    def test_update_unknown_id_404(self, admin):
        for res, body in [("teachers", {"name": "X"}), ("classes", {"name": "X"}), ("subjects", {"name": "X"}), ("students", {"name": "X", "class_name": "7A"})]:
            assert admin.put(f"{BASE_URL}/api/{res}/nope", json=body, timeout=30).status_code == 404, res
            assert admin.delete(f"{BASE_URL}/api/{res}/nope", timeout=30).status_code == 404, res

    def test_empty_name_rejected(self, admin):
        assert admin.post(f"{BASE_URL}/api/classes", json={"name": ""}, timeout=30).status_code == 422

    def test_guru_forbidden_on_master_writes(self, guru, admin):
        # create with admin so the target exists
        cid = admin.post(f"{BASE_URL}/api/classes", json={"name": "TEST_RBAC"}, timeout=30).json()["id"]
        tid = admin.post(f"{BASE_URL}/api/teachers", json={"name": "TEST_RBAC"}, timeout=30).json()["id"]
        sid = admin.post(f"{BASE_URL}/api/subjects", json={"name": "TEST_RBAC"}, timeout=30).json()["id"]
        stid = admin.post(f"{BASE_URL}/api/students", json={"name": "TEST_RBAC", "class_name": "7A"}, timeout=30).json()["id"]
        checks = [("classes", cid, {"name": "hack"}), ("teachers", tid, {"name": "hack"}), ("subjects", sid, {"name": "hack"}), ("students", stid, {"name": "hack", "class_name": "7A"})]
        for res, rid, body in checks:
            assert guru.put(f"{BASE_URL}/api/{res}/{rid}", json=body, timeout=30).status_code == 403, f"PUT {res}"
            assert guru.delete(f"{BASE_URL}/api/{res}/{rid}", timeout=30).status_code == 403, f"DELETE {res}"
            assert guru.post(f"{BASE_URL}/api/{res}", json=body, timeout=30).status_code == 403, f"POST {res}"
        for res, rid, _ in checks:
            admin.delete(f"{BASE_URL}/api/{res}/{rid}", timeout=30)


# ---------- Dashboard stats ----------
class TestDashboardStats:
    def test_stats_reflect_real_data(self, admin):
        r = admin.get(f"{BASE_URL}/api/dashboard/stats", timeout=30)
        assert r.status_code == 200, r.text[:300]
        s = r.json()
        masters = admin.get(f"{BASE_URL}/api/masters", timeout=30).json()
        journals = admin.get(f"{BASE_URL}/api/journals", timeout=30).json()
        assert s["classes_count"] == len(masters["classes"])
        assert s["students_count"] == len(masters["students"])
        assert s["journals_count"] == len(journals)
        assert 0 <= s["attendance_rate"] <= 100
        assert isinstance(s["top_absent"], list) and len(s["top_absent"]) <= 3
        assert s["today_day"] in ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]

    def test_stats_change_when_class_added(self, admin):
        before = admin.get(f"{BASE_URL}/api/dashboard/stats", timeout=30).json()["classes_count"]
        cid = admin.post(f"{BASE_URL}/api/classes", json={"name": "TEST_STAT"}, timeout=30).json()["id"]
        after = admin.get(f"{BASE_URL}/api/dashboard/stats", timeout=30).json()["classes_count"]
        admin.delete(f"{BASE_URL}/api/classes/{cid}", timeout=30)
        assert after == before + 1

    def test_today_schedule_matches_today_day(self, admin):
        s = admin.get(f"{BASE_URL}/api/dashboard/stats", timeout=30).json()
        for item in s["today_schedule"]:
            assert item["day"] == s["today_day"]

    def test_absent_stats_from_real_attendance(self, admin):
        import datetime
        month = datetime.datetime.utcnow().strftime("%Y-%m")
        date = f"{month}-15"
        payload = {"date": date, "class_name": "7A", "subject": "Matematika", "entries": [
            {"student": "TEST_Alpa Anak", "status": "A", "note": ""},
            {"student": "Alya Putri", "status": "H", "note": ""},
        ]}
        assert admin.post(f"{BASE_URL}/api/attendance", json=payload, timeout=30).status_code == 200
        s = admin.get(f"{BASE_URL}/api/dashboard/stats", timeout=30).json()
        names = [x["name"] for x in s["top_absent"]]
        assert "TEST_Alpa Anak" in names, f"absent student missing from top_absent: {s['top_absent']}"
        assert s["attendance_rate"] > 0


# ---------- Reports filters ----------
class TestReportFilters:
    def test_rows_filter_by_class(self, admin):
        all_rows = admin.get(f"{BASE_URL}/api/reports/rows", params={"kind": "attendance"}, timeout=30)
        assert all_rows.status_code == 200
        all_rows = all_rows.json()
        filtered = admin.get(f"{BASE_URL}/api/reports/rows", params={"kind": "attendance", "class_name": "7A"}, timeout=30).json()
        assert len(filtered) <= len(all_rows)
        assert all(r["class_name"] == "7A" for r in filtered), set(r["class_name"] for r in filtered)

    def test_rows_filter_combo(self, admin):
        import datetime
        month = datetime.datetime.utcnow().strftime("%Y-%m")
        rows = admin.get(f"{BASE_URL}/api/reports/rows", params={"kind": "attendance", "class_name": "7A", "subject": "Matematika", "period": month}, timeout=30).json()
        for r in rows:
            assert r["class_name"] == "7A" and r["subject"] == "Matematika" and r["date"].startswith(month), r

    def test_rows_numbering_sequential(self, admin):
        rows = admin.get(f"{BASE_URL}/api/reports/rows", params={"kind": "attendance", "class_name": "7A"}, timeout=30).json()
        assert [r["no"] for r in rows] == list(range(1, len(rows) + 1))

    def test_rows_kinds(self, admin):
        for kind in ["attendance", "grades", "journals"]:
            r = admin.get(f"{BASE_URL}/api/reports/rows", params={"kind": kind}, timeout=30)
            assert r.status_code == 200, f"{kind} -> {r.status_code} {r.text[:200]}"
            assert isinstance(r.json(), list)

    def test_export_excel_respects_filters(self, admin):
        r_all = admin.get(f"{BASE_URL}/api/reports/export", params={"kind": "attendance"}, timeout=60)
        r_f = admin.get(f"{BASE_URL}/api/reports/export", params={"kind": "attendance", "class_name": "7A"}, timeout=60)
        assert r_all.status_code == 200 and r_f.status_code == 200
        assert "spreadsheetml" in r_all.headers["content-type"]
        wb_all = load_workbook(io.BytesIO(r_all.content)).active
        wb_f = load_workbook(io.BytesIO(r_f.content)).active
        rows_all = list(wb_all.iter_rows(min_row=2, values_only=True))
        rows_f = list(wb_f.iter_rows(min_row=2, values_only=True))
        assert len(rows_f) <= len(rows_all)
        assert all(row[2] == "7A" for row in rows_f), set(row[2] for row in rows_f)

    def test_print_html_respects_filters(self, admin):
        r = admin.get(f"{BASE_URL}/api/reports/print", params={"kind": "attendance", "class_name": "7A", "period": "2026-06"}, timeout=60)
        assert r.status_code == 200
        html = r.text
        assert "Kelas: 7A" in html and "Periode: 2026-06" in html
        for other in ["7B", "8A", "9A"]:
            assert f"<td>{other}</td>" not in html, f"unfiltered class {other} present in print output"

    def test_export_empty_filter_still_valid(self, admin):
        r = admin.get(f"{BASE_URL}/api/reports/export", params={"kind": "attendance", "class_name": "NO_SUCH_CLASS"}, timeout=60)
        assert r.status_code == 200
        sheet = load_workbook(io.BytesIO(r.content)).active
        assert list(sheet.iter_rows(min_row=2, values_only=True)) == []


# ---------- Regression: attendance / grades / journals ----------
class TestRegression:
    def test_attendance_upsert_and_read(self, guru):
        body = {"date": "2026-07-02", "class_name": "7A", "subject": "IPA", "entries": [{"student": "Alya Putri", "status": "S", "note": "TEST"}]}
        assert guru.post(f"{BASE_URL}/api/attendance", json=body, timeout=30).status_code == 200
        got = guru.get(f"{BASE_URL}/api/attendance", params={"date": "2026-07-02", "class_name": "7A", "subject": "IPA"}, timeout=30).json()
        assert got["entries"][0]["status"] == "S" and "_id" not in got

    def test_grades_save_and_read(self, guru):
        body = {"date": "2026-07-02", "class_name": "7A", "subject": "IPA", "assessment_type": "Formatif", "entries": [{"student": "Alya Putri", "score": 88}]}
        assert guru.post(f"{BASE_URL}/api/grades", json=body, timeout=30).status_code == 200
        got = guru.get(f"{BASE_URL}/api/grades", params={"date": "2026-07-02", "class_name": "7A", "subject": "IPA", "assessment_type": "Formatif"}, timeout=30).json()
        assert got["entries"][0]["score"] == 88

    def test_grade_out_of_range_rejected(self, guru):
        body = {"date": "2026-07-02", "class_name": "7A", "subject": "IPA", "assessment_type": "Formatif", "entries": [{"student": "X", "score": 150}]}
        assert guru.post(f"{BASE_URL}/api/grades", json=body, timeout=30).status_code == 422

    def test_journal_create_and_list(self, guru):
        body = {"date": "2026-07-02", "period": 3, "class_name": "7A", "subject": "IPA", "topic": "TEST_Topik", "activity": "TEST_Kegiatan", "reflection": ""}
        r = guru.post(f"{BASE_URL}/api/journals", json=body, timeout=30)
        assert r.status_code == 200
        assert "_id" not in r.json()
        lst = guru.get(f"{BASE_URL}/api/journals", timeout=30).json()
        assert any(j["topic"] == "TEST_Topik" for j in lst)

    def test_settings_guru_forbidden(self, guru):
        r = guru.put(f"{BASE_URL}/api/settings", json={"school": "Hack", "address": "x", "principal": "y", "nip": "1"}, timeout=30)
        assert r.status_code == 403
