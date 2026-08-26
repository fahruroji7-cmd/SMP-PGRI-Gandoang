"""AbsenSPG backend regression suite (iteration 6): auth, masters, attendance, grades, journals, settings, reports."""
import io
import os
import re
import uuid
from pathlib import Path

import openpyxl
import pytest
import requests
from dotenv import dotenv_values

frontend_env = dotenv_values("/app/frontend/.env")
base_url = os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")
if not base_url:
    raise RuntimeError("REACT_APP_BACKEND_URL missing")
BASE_URL = base_url.rstrip("/")

TAG = uuid.uuid4().hex[:6]


def _creds():
    content = Path("/app/memory/test_credentials.md").read_text(encoding="utf-8")
    pairs = re.findall(r"Email:\s*`([^`]+)`\s*\n-\s*Password:\s*`([^`]+)`", content)
    return pairs


CREDS = dict((e, p) for e, p in _creds())


def login(email):
    session = requests.Session()
    response = session.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": CREDS[email]}, timeout=30)
    if response.status_code != 200:
        pytest.fail(f"login failed for {email}: {response.status_code} {response.text[:200]}")
    return session


@pytest.fixture(scope="session")
def admin():
    return login("admin@absenspg.local")


@pytest.fixture(scope="session")
def guru():
    return login("guru@absenspg.local")


# ---------- auth ----------
class TestAuth:
    def test_root(self):
        r = requests.get(f"{BASE_URL}/api/", timeout=30)
        assert r.status_code == 200 and "online" in r.json()["message"]

    @pytest.mark.parametrize("email,role", [("admin@absenspg.local", "Admin"), ("guru@absenspg.local", "Guru")])
    def test_login_sets_httponly_cookie(self, email, role):
        s = requests.Session()
        r = s.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": CREDS[email]}, timeout=30)
        assert r.status_code == 200
        body = r.json()
        assert body["role"] == role and body["email"] == email and body["id"]
        assert "password_hash" not in body and "_id" not in body
        cookie_header = r.headers.get("set-cookie", "")
        assert "access_token=" in cookie_header
        assert "HttpOnly" in cookie_header and "Secure" in cookie_header
        me = s.get(f"{BASE_URL}/api/auth/me", timeout=30)
        assert me.status_code == 200 and me.json()["role"] == role
        perms = s.get(f"{BASE_URL}/api/auth/permissions", timeout=30)
        assert perms.status_code == 200
        assert perms.json()["can_manage_master_data"] is (role == "Admin")

    def test_login_bad_password(self):
        r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": "admin@absenspg.local", "password": "wrong"}, timeout=30)
        assert r.status_code == 401 and "salah" in r.json()["detail"]

    def test_login_unknown_email(self):
        r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": f"nobody{TAG}@x.local", "password": "Whatever1!"}, timeout=30)
        assert r.status_code == 401

    def test_login_validation_empty_password(self):
        r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": "admin@absenspg.local", "password": ""}, timeout=30)
        assert r.status_code == 422

    def test_bcrypt_hash_format(self):
        # bcrypt hash must be $2b$ format in DB
        import subprocess
        out = subprocess.run(
            ["python", "-c",
             "import asyncio,os;from dotenv import load_dotenv;load_dotenv('/app/backend/.env');"
             "from motor.motor_asyncio import AsyncIOMotorClient;"
             "c=AsyncIOMotorClient(os.environ['MONGO_URL']);d=c[os.environ['DB_NAME']];"
             "print(asyncio.get_event_loop().run_until_complete(d.users.find_one({'email':'admin@absenspg.local'}))['password_hash'])"],
            capture_output=True, text=True, timeout=60)
        assert out.stdout.strip().startswith("$2b$"), out.stdout + out.stderr

    def test_logout(self):
        s = login("guru@absenspg.local")
        assert s.post(f"{BASE_URL}/api/auth/logout", timeout=30).status_code == 200
        assert s.get(f"{BASE_URL}/api/auth/me", timeout=30).status_code == 401

    def test_invalid_token_rejected(self):
        r = requests.get(f"{BASE_URL}/api/auth/me", headers={"Authorization": "Bearer bogus.token.here"}, timeout=30)
        assert r.status_code == 401

    @pytest.mark.parametrize("path", ["/api/masters", "/api/students", "/api/journals", "/api/settings", "/api/reports/rows", "/api/reports/export", "/api/reports/print"])
    def test_unauthenticated_401(self, path):
        assert requests.get(f"{BASE_URL}{path}", timeout=30).status_code == 401

    def test_cors_allows_credentials_with_explicit_origin(self):
        # NOTE: the k8s ingress/CDN answers OPTIONS with "*" before FastAPI, so CORS
        # middleware config is asserted against the app directly.
        origin = BASE_URL
        r = requests.options("http://localhost:8001/api/auth/login", headers={"Origin": origin, "Access-Control-Request-Method": "POST"}, timeout=30)
        assert r.headers.get("access-control-allow-credentials") == "true"
        assert r.headers.get("access-control-allow-origin") == origin


# ---------- masters + role guards ----------
class TestMasters:
    created = {"students": [], "classes": [], "subjects": [], "teachers": []}

    def test_masters_shape(self, guru):
        r = guru.get(f"{BASE_URL}/api/masters", timeout=30)
        assert r.status_code == 200
        data = r.json()
        for key in ["classes", "subjects", "teachers", "students"]:
            assert isinstance(data[key], list) and data[key]
            assert all("_id" not in row for row in data[key])
        names = [c["name"] for c in data["classes"]]
        assert {"7A", "7B", "8A", "9A"} <= set(names)

    @pytest.mark.parametrize("class_name,count", [("7A", 6), ("7B", 3), ("8A", 3)])
    def test_students_by_class(self, guru, class_name, count):
        r = guru.get(f"{BASE_URL}/api/students", params={"class_name": class_name}, timeout=30)
        assert r.status_code == 200
        rows = r.json()
        assert all(s["class_name"] == class_name for s in rows)
        assert len(rows) >= count

    @pytest.mark.parametrize("endpoint", ["students", "classes", "subjects", "teachers"])
    def test_guru_forbidden(self, guru, endpoint):
        payload = {"name": f"TEST_{TAG}", "class_name": "7A"} if endpoint == "students" else {"name": f"TEST_{TAG}"}
        r = guru.post(f"{BASE_URL}/api/{endpoint}", json=payload, timeout=30)
        assert r.status_code == 403

    @pytest.mark.parametrize("endpoint,key", [("students", "students"), ("classes", "classes"), ("subjects", "subjects"), ("teachers", "teachers")])
    def test_admin_create_and_persist(self, admin, endpoint, key):
        name = f"TEST_{endpoint}_{TAG}"
        payload = {"name": name, "class_name": "7A"} if endpoint == "students" else {"name": name}
        r = admin.post(f"{BASE_URL}/api/{endpoint}", json=payload, timeout=30)
        assert r.status_code == 200, r.text
        doc = r.json()
        assert doc["name"] == name and "_id" not in doc and isinstance(doc["id"], str)
        self.created[endpoint].append(doc["id"])
        listing = admin.get(f"{BASE_URL}/api/masters", timeout=30).json()[key]
        assert any(row["id"] == doc["id"] and row["name"] == name for row in listing)

    def test_student_validation_empty_name(self, admin):
        r = admin.post(f"{BASE_URL}/api/students", json={"name": "", "class_name": "7A"}, timeout=30)
        assert r.status_code == 422


# ---------- attendance ----------
class TestAttendance:
    date = "2026-07-02"
    payload = {
        "date": date, "class_name": "7B", "subject": "IPA",
        "entries": [
            {"student": "Gita Ayu", "status": "S", "note": "TEST_demam"},
            {"student": "Hadi Saputra", "status": "H", "note": ""},
            {"student": "Intan Permata", "status": "A", "note": "TEST_tanpa keterangan"},
        ],
    }

    def test_save_and_roundtrip(self, guru):
        r = guru.post(f"{BASE_URL}/api/attendance", json=self.payload, timeout=30)
        assert r.status_code == 200 and r.json()["count"] == 3
        got = guru.get(f"{BASE_URL}/api/attendance", params={"date": self.date, "class_name": "7B", "subject": "IPA"}, timeout=30)
        assert got.status_code == 200
        data = got.json()
        assert "_id" not in data
        entries = {e["student"]: e for e in data["entries"]}
        assert entries["Gita Ayu"]["status"] == "S" and entries["Gita Ayu"]["note"] == "TEST_demam"
        assert entries["Intan Permata"]["status"] == "A"
        assert data["recorded_by"]

    def test_upsert_overwrites_not_duplicates(self, guru):
        changed = dict(self.payload)
        changed["entries"] = [{"student": "Gita Ayu", "status": "H", "note": ""}]
        assert guru.post(f"{BASE_URL}/api/attendance", json=changed, timeout=30).status_code == 200
        data = guru.get(f"{BASE_URL}/api/attendance", params={"date": self.date, "class_name": "7B", "subject": "IPA"}, timeout=30).json()
        assert len(data["entries"]) == 1 and data["entries"][0]["status"] == "H"
        # restore multi-entry state for report tests
        assert guru.post(f"{BASE_URL}/api/attendance", json=self.payload, timeout=30).status_code == 200

    def test_empty_result_for_unknown_key(self, guru):
        data = guru.get(f"{BASE_URL}/api/attendance", params={"date": "1999-01-01", "class_name": "7A", "subject": "IPA"}, timeout=30).json()
        assert data["entries"] == []

    def test_invalid_status_rejected(self, guru):
        bad = dict(self.payload)
        bad["entries"] = [{"student": "Gita Ayu", "status": "X"}]
        assert guru.post(f"{BASE_URL}/api/attendance", json=bad, timeout=30).status_code == 422

    def test_missing_params_on_get(self, guru):
        assert guru.get(f"{BASE_URL}/api/attendance", params={"date": self.date}, timeout=30).status_code == 422


# ---------- grades ----------
class TestGrades:
    date = "2026-07-02"
    payload = {
        "date": date, "class_name": "8A", "subject": "Informatika", "assessment_type": "Formatif",
        "entries": [
            {"student": "Joko Santoso", "score": 92},
            {"student": "Kirana Dewi", "score": 78},
            {"student": "Lukman Hakim", "score": 60},
        ],
    }

    def test_save_and_roundtrip(self, guru):
        r = guru.post(f"{BASE_URL}/api/grades", json=self.payload, timeout=30)
        assert r.status_code == 200 and r.json()["count"] == 3
        data = guru.get(f"{BASE_URL}/api/grades", params={"date": self.date, "class_name": "8A", "subject": "Informatika", "assessment_type": "Formatif"}, timeout=30).json()
        scores = {e["student"]: e["score"] for e in data["entries"]}
        assert scores["Joko Santoso"] == 92 and scores["Lukman Hakim"] == 60

    def test_assessment_type_is_part_of_key(self, guru):
        other = dict(self.payload)
        other["assessment_type"] = "Sumatif Lingkup Materi"
        other["entries"] = [{"student": "Joko Santoso", "score": 55}]
        assert guru.post(f"{BASE_URL}/api/grades", json=other, timeout=30).status_code == 200
        formatif = guru.get(f"{BASE_URL}/api/grades", params={"date": self.date, "class_name": "8A", "subject": "Informatika", "assessment_type": "Formatif"}, timeout=30).json()
        assert {e["student"]: e["score"] for e in formatif["entries"]}["Joko Santoso"] == 92

    def test_score_out_of_range_rejected(self, guru):
        bad = dict(self.payload)
        bad["entries"] = [{"student": "Joko Santoso", "score": 150}]
        assert guru.post(f"{BASE_URL}/api/grades", json=bad, timeout=30).status_code == 422


# ---------- journals ----------
class TestJournals:
    def test_create_and_list(self, guru):
        payload = {"date": "2026-07-02", "period": 3, "class_name": "7A", "subject": "Matematika",
                   "topic": f"TEST_topik_{TAG}", "activity": "TEST_kegiatan inti", "reflection": ""}
        r = guru.post(f"{BASE_URL}/api/journals", json=payload, timeout=30)
        assert r.status_code == 200
        doc = r.json()
        assert doc["topic"] == payload["topic"] and doc["teacher"] and "_id" not in doc
        rows = guru.get(f"{BASE_URL}/api/journals", timeout=30).json()
        assert rows[0]["id"] == doc["id"], "newest journal should be first"
        assert any(x["topic"] == payload["topic"] for x in rows)

    def test_empty_topic_rejected(self, guru):
        r = guru.post(f"{BASE_URL}/api/journals", json={"date": "2026-07-02", "period": 1, "class_name": "7A", "subject": "Matematika", "topic": "", "activity": "x"}, timeout=30)
        assert r.status_code == 422


# ---------- settings ----------
class TestSettings:
    def test_guru_cannot_update(self, guru):
        r = guru.put(f"{BASE_URL}/api/settings", json={"school": "HACK", "address": "a", "principal": "p", "nip": "1"}, timeout=30)
        assert r.status_code == 403

    def test_admin_update_and_restore(self, admin):
        original = admin.get(f"{BASE_URL}/api/settings", timeout=30).json()
        assert "_id" not in original
        new = {"school": f"TEST_Sekolah_{TAG}", "address": original["address"], "principal": original["principal"], "nip": original["nip"]}
        assert admin.put(f"{BASE_URL}/api/settings", json=new, timeout=30).status_code == 200
        assert admin.get(f"{BASE_URL}/api/settings", timeout=30).json()["school"] == new["school"]
        # print report should use saved school name
        html = admin.get(f"{BASE_URL}/api/reports/print", params={"kind": "journals"}, timeout=30).text
        assert new["school"] in html
        restore = {k: original[k] for k in ["school", "address", "principal", "nip"]}
        assert admin.put(f"{BASE_URL}/api/settings", json=restore, timeout=30).status_code == 200
        assert admin.get(f"{BASE_URL}/api/settings", timeout=30).json()["school"] == original["school"]


# ---------- reports from real DB ----------
class TestReports:
    def test_attendance_rows_from_db(self, guru):
        rows = guru.get(f"{BASE_URL}/api/reports/rows", params={"kind": "attendance"}, timeout=30).json()
        assert rows and rows[0]["no"] == 1
        target = [r for r in rows if r["student"] == "Gita Ayu" and r["date"] == "2026-07-02" and r["subject"] == "IPA"]
        assert target and target[0]["status"] == "Sakit", target
        assert all("_id" not in r for r in rows)

    def test_grades_rows_predicate(self, guru):
        rows = guru.get(f"{BASE_URL}/api/reports/rows", params={"kind": "grades"}, timeout=30).json()
        subset = [r for r in rows if r["date"] == "2026-07-02" and r["subject"] == "Informatika"]
        # NOTE: report rows do not carry assessment_type, so both Formatif and Sumatif
        # entries for the same student/date/subject appear identical apart from value.
        joko = [r for r in subset if r["student"] == "Joko Santoso"]
        assert any(r["value"] == "92" and r["status"] == "Sangat baik" for r in joko), joko
        lukman = [r for r in subset if r["student"] == "Lukman Hakim"]
        assert lukman and lukman[0]["status"] == "Perlu bimbingan"

    def test_journal_rows(self, guru):
        # create own journal first: xdist runs classes in parallel, so do not rely on other tests
        guru.post(f"{BASE_URL}/api/journals", json={"date": "2026-07-02", "period": 2, "class_name": "7A", "subject": "IPA",
                                                    "topic": f"TEST_topik_{TAG}", "activity": "TEST_kegiatan", "reflection": ""}, timeout=30)
        rows = guru.get(f"{BASE_URL}/api/reports/rows", params={"kind": "journals"}, timeout=30).json()
        assert rows and all(r["status"] == "Tersimpan" for r in rows)
        assert any(TAG in r["value"] for r in rows)

    @pytest.mark.parametrize("kind", ["attendance", "grades", "journals"])
    def test_export_xlsx_matches_rows(self, admin, kind):
        rows = admin.get(f"{BASE_URL}/api/reports/rows", params={"kind": kind}, timeout=30).json()
        r = admin.get(f"{BASE_URL}/api/reports/export", params={"kind": kind}, timeout=60)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        assert f"absenspg-{kind}.xlsx" in r.headers.get("content-disposition", "")
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
        sheet = wb.active
        assert sheet.cell(1, 1).value == "No" and sheet.cell(1, 7).value == "Nilai / Materi"
        assert sheet.max_row == len(rows) + 1, f"{kind}: xlsx rows {sheet.max_row} vs api {len(rows)}"

    @pytest.mark.parametrize("kind,title", [("attendance", "Rekap Absensi Siswa"), ("grades", "Rekap Nilai Siswa"), ("journals", "Rekap Jurnal Mengajar")])
    def test_print_html_title_and_content(self, admin, kind, title):
        r = admin.get(f"{BASE_URL}/api/reports/print", params={"kind": kind}, timeout=60)
        assert r.status_code == 200 and "text/html" in r.headers.get("content-type", "")
        assert f"<title>{title}</title>" in r.text
        assert "window.print()" in r.text and "<table>" in r.text
        rows = admin.get(f"{BASE_URL}/api/reports/rows", params={"kind": kind}, timeout=30).json()
        assert r.text.count("<tr>") == len(rows) + 1

    def test_unknown_kind_falls_back_to_attendance(self, guru):
        rows = guru.get(f"{BASE_URL}/api/reports/rows", params={"kind": "bogus"}, timeout=30).json()
        attendance = guru.get(f"{BASE_URL}/api/reports/rows", params={"kind": "attendance"}, timeout=30).json()
        assert rows == attendance
